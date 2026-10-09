"""Read Excel (.xlsx) and CSV files into a header row + data rows."""
import csv
import io
import os
import re

from openpyxl import Workbook, load_workbook

from .merge import valid_email


class SheetError(Exception):
    pass


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def sheet_names(path: str) -> list:
    if path.lower().endswith((".csv", ".txt", ".tsv")):
        return ["CSV"]
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        raise SheetError("Couldn't open that file. Please use an Excel .xlsx file or a .csv "
                         "(in Excel: File > Save As > Excel Workbook).")
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def read(path: str, sheet: str = None):
    """Returns (headers, rows). Rows are lists of strings padded to len(headers)."""
    if path.lower().endswith((".csv", ".txt", ".tsv")):
        raw = open(path, "rb").read().decode("utf-8-sig", errors="replace")
        first = raw.splitlines()[0] if raw.strip() else ""
        delim = max([",", ";", "\t"], key=lambda d: first.count(d))
        table = [[c.strip() for c in r] for r in csv.reader(io.StringIO(raw), delimiter=delim)]
    else:
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
        except Exception:
            raise SheetError("Couldn't open that file. Please use an Excel .xlsx file or a .csv.")
        try:
            ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb[wb.sheetnames[0]]
            table = [[_cell(c) for c in r] for r in ws.iter_rows(values_only=True)]
        finally:
            wb.close()
    table = [r for r in table if any(c for c in r)]
    if not table:
        return [], []
    width = max(len(r) for r in table)
    # drop trailing completely-empty columns
    while width and not any((r[width - 1] if len(r) >= width else "") for r in table):
        width -= 1
    headers, seen = [], set()
    for i in range(width):
        h = (table[0][i] if i < len(table[0]) else "").strip() or f"Column {i + 1}"
        base, n = h, 2
        while h.lower() in seen:
            h, n = f"{base} ({n})", n + 1
        seen.add(h.lower())
        headers.append(h)
    data = [(r + [""] * width)[:width] for r in table[1:]]
    return headers, data


def detect_email_column(headers, rows) -> str:
    norm = [re.sub(r"[\s_\-]", "", h.lower()) for h in headers]
    for i, n in enumerate(norm):
        if n in ("email", "emailaddress", "emailaddr", "mail", "e-mail"):
            return headers[i]
    for i, n in enumerate(norm):
        if "email" in n or "mail" in n:
            return headers[i]
    best, best_score = "", 0.0
    sample = rows[:50]
    for i, h in enumerate(headers):
        vals = [r[i] for r in sample if r[i]]
        if vals:
            score = sum(valid_email(v) for v in vals) / len(sample)
            if score > best_score and score >= 0.5:
                best, best_score = h, score
    return best


def write_xlsx(path: str, columns: list, records: list):
    wb = Workbook()
    ws = wb.active
    ws.title = "Contacts"
    ws.append(columns)
    for rec in records:
        ws.append([rec.get(c, "") for c in columns])
    for i, c in enumerate(columns, 1):
        width = max([len(str(c))] + [len(str(r.get(c, ""))) for r in records[:200]]) + 2
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(width, 50)
    wb.save(path)
