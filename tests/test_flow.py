import email
import email.policy
import glob
import io
import os
import tempfile
import time

TMP = tempfile.mkdtemp(prefix="mm-test-")
os.environ["DATA_DIR"] = TMP

import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import Workbook  # noqa: E402

from app import db, worker  # noqa: E402
from app.main import app  # noqa: E402


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        yield c


def xlsx_bytes(rows, sheet="Clients"):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    for r in rows:
        ws.append(r)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def wait_done(client, cid, timeout=30):
    end = time.time() + timeout
    while time.time() < end:
        c = client.get(f"/api/campaigns/{cid}").json()
        if c["status"] in ("done", "paused"):
            return c
        time.sleep(0.2)
    raise AssertionError("campaign did not finish")


def eml(folder):
    out = {}
    for f in glob.glob(os.path.join(TMP, "outbox", "*", folder, "*.eml")):
        m = email.message_from_binary_file(open(f, "rb"), policy=email.policy.default)
        out[m["To"]] = m
    return out


def text_of(m):
    part = m.get_body(preferencelist=("plain",))
    return part.get_content()


S = {}


def test_setup_account(client):
    r = client.post("/api/profiles", json={"name": "Test", "kind": "dryrun", "from_email": "wife@work.test", "per_minute": 600})
    assert r.status_code == 200
    S["pid"] = r.json()["id"]
    assert client.post(f"/api/profiles/{S['pid']}/check").json()["ok"]


def test_import_review_and_commit(client):
    data = xlsx_bytes([
        ["First Name", "Email Address", "Company", "Suburb"],
        ["Ann", "ann@example.com", "Acme", "Hutt"],
        ["Bob", "BOB@example.com ", "Beta", ""],
        ["Cy", "not-an-email", "Gamma", "Nelson"],
        ["Dup", "ann@example.com", "Acme2", "X"],
        ["NoMail", "", "Delta", "Y"],
        ["Eve", "eve@example.com", "Echo", "Wellington"],
    ])
    r = client.post("/api/import/preview", files={"file": ("clients.xlsx", data)})
    assert r.status_code == 200, r.text
    pv = r.json()
    assert pv["email_column"] == "Email Address"
    assert pv["headers"] == ["First Name", "Email Address", "Company", "Suburb"]
    assert pv["sheets"] == ["Clients"] and pv["row_count"] == 6
    r = client.post(f"/api/import/{pv['token']}/commit", json={
        "sheet": "Clients", "email_column": "Email Address", "new_list_name": "Clients", "update_existing": False}).json()
    assert r["added"] == 4 and r["duplicates"] == 1 and r["no_email"] == 1 and r["invalid"] == 1
    S["lid"] = r["list_id"]
    d = client.get(f"/api/lists/{S['lid']}/contacts").json()
    assert d["columns"] == ["Email", "First Name", "Company", "Suburb"]
    emails = {c["email"]: c for c in d["contacts"]}
    assert "bob@example.com" in emails                       # trimmed + lower-cased
    assert emails["not-an-email"]["valid"] is False         # flagged, kept so it can be fixed
    S["ids"] = {e: c["id"] for e, c in emails.items()}


def test_edit_add_column_and_fix_email(client):
    ids = S["ids"]
    d = client.get(f"/api/lists/{S['lid']}/contacts").json()
    cy = next(c for c in d["contacts"] if c["email"] == "not-an-email")
    r = client.put(f"/api/contacts/{cy['id']}", json={"data": {**cy["data"], "Email": "cy@example.com"}})
    assert r.status_code == 200 and r.json()["valid"]
    # duplicate email on edit is rejected
    r = client.put(f"/api/contacts/{cy['id']}", json={"data": {**cy["data"], "Email": "ann@example.com"}})
    assert r.status_code == 409
    assert client.post(f"/api/lists/{S['lid']}/columns", json={"name": "Nickname"}).status_code == 200
    assert client.post(f"/api/lists/{S['lid']}/columns", json={"name": "nickname"}).status_code == 409
    r = client.post(f"/api/lists/{S['lid']}/contacts", json={"data": {"Email": "fay@example.com", "First Name": "Fay <b>"}})
    assert r.status_code == 200
    S["fay"] = r.json()["id"]
    # 'don't email' flag on Eve
    assert client.put(f"/api/contacts/{ids['eve@example.com']}", json={"skip": True}).json()["skip"] is True


def test_export_roundtrip(client):
    r = client.get(f"/api/lists/{S['lid']}/export")
    assert r.status_code == 200 and r.content[:2] == b"PK"
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(r.content)).active
    assert [c.value for c in ws[1]][:3] == ["Email", "First Name", "Company"]


def test_preview_and_unknown_fields(client):
    r = client.post("/api/campaigns", json={"name": "Spring"})
    cid = S["cid"] = r.json()["id"]
    assert r.json()["profile_id"] == S["pid"] and r.json()["list_id"] == S["lid"]
    client.put(f"/api/campaigns/{cid}", json={
        "subject": "Hi {{First Name|there}} at {{Company}}",
        "body": "Dear {{ First Name }},\nSuburb: {{Suburb|unknown}}\n{{Missing}}\n"})
    pv = client.post(f"/api/campaigns/{cid}/preview", json={}).json()
    assert pv["subject"] == "Hi Ann at Acme"
    assert pv["unknown_fields"] == ["Missing"]
    assert "Suburb: Hutt" in pv["body"]


def test_attachment_upload(client):
    r = client.post(f"/api/campaigns/{S['cid']}/attachments", files={"file": ("price list.txt", b"hello attachment")})
    assert r.status_code == 200
    assert client.get(f"/api/campaigns/{S['cid']}").json()["attachments"][0]["filename"] == "price list.txt"


def test_suppression(client):
    client.post("/api/suppressions", json={"emails": "bob@example.com"})
    assert client.get("/api/suppressions").json()[0]["email"] == "bob@example.com"


def test_create_drafts_then_send(client):
    cid = S["cid"]
    plan = client.get(f"/api/campaigns/{cid}/plan").json()
    # ann, cy, fay eligible; bob opted out; eve flagged
    assert plan["to_draft"] == 3 and plan["skipped"] == 2 and plan["supports_drafts"]
    assert client.post(f"/api/campaigns/{cid}/run", json={"mode": "draft"}).status_code == 200
    c = wait_done(client, cid)
    assert c["status"] == "done" and c["counts"]["drafted"] == 3 and c["counts"]["skipped"] == 2
    drafts = eml("drafts")
    assert set(drafts) == {"ann@example.com", "cy@example.com", "fay@example.com"}
    assert not eml("sent")                                   # nothing sent yet
    ann = drafts["ann@example.com"]
    assert ann["Subject"] == "Hi Ann at Acme" and ann["From"] == "wife@work.test"
    assert "Suburb: Hutt" in text_of(ann)
    assert [p.get_filename() for p in ann.iter_attachments()] == ["price list.txt"]

    # now send: the existing drafts are the things that get sent (no second copies)
    assert client.post(f"/api/campaigns/{cid}/run", json={"mode": "send"}).status_code == 200
    c = wait_done(client, cid)
    assert c["counts"]["sent"] == 3
    assert set(eml("sent")) == {"ann@example.com", "cy@example.com", "fay@example.com"}
    assert not eml("drafts")                                 # drafts moved to Sent, not duplicated


def test_rerun_only_new_people(client):
    cid = S["cid"]
    client.put(f"/api/contacts/{S['ids']['eve@example.com']}", json={"skip": False})
    client.post(f"/api/lists/{S['lid']}/contacts", json={"data": {"Email": "gus@example.com", "First Name": "Gus"}})
    plan = client.get(f"/api/campaigns/{cid}/plan").json()
    assert plan["already_sent"] == 3 and plan["to_send"] == 2   # eve + gus only
    client.post(f"/api/campaigns/{cid}/run", json={"mode": "send"})
    c = wait_done(client, cid)
    assert c["counts"]["sent"] == 5
    assert len(eml("sent")) == 5


def test_html_mode_escapes_values(client):
    r = client.post("/api/campaigns", json={"name": "Styled"}).json()
    client.put(f"/api/campaigns/{r['id']}", json={"subject": "News", "is_html": True,
                                                  "body": "<p>Hello <b>{{First Name}}</b></p>"})
    pv = client.post(f"/api/campaigns/{r['id']}/preview", json={"contact_id": S["fay"]}).json()
    assert pv["body"] == "<p>Hello <b>Fay &lt;b&gt;</b></p>"      # contact data can't inject markup
    assert client.post(f"/api/campaigns/{r['id']}/test", json={"to": "me@work.test"}).status_code == 200
    m = eml("sent")["me@work.test"]
    assert m.get_body(preferencelist=("html",)).get_content().startswith("<p>Hello")
    assert "Hello" in text_of(m)                              # plain-text alternative included


def test_crash_mid_send_is_never_resent(client):
    with db.connect() as c:
        cid = S["cid"]
        mid = db.row(c, "SELECT id FROM messages WHERE campaign_id=? AND email='gus@example.com'", cid)["id"]
        c.execute("UPDATE messages SET status='sending' WHERE id=?", (mid,))
    worker.start()   # what happens on app restart
    time.sleep(0.5)
    with db.connect() as c:
        assert db.row(c, "SELECT status FROM messages WHERE id=?", mid)["status"] == "uncertain"
    before = len(eml("sent"))
    client.put(f"/api/campaigns/{cid}", json={"name": "Spring"})
    client.post(f"/api/campaigns/{cid}/run", json={"mode": "send"})
    wait_done(client, cid)
    assert len(eml("sent")) == before                        # still not re-sent
    assert client.post(f"/api/messages/{mid}/mark", json={"status": "sent"}).status_code == 200


def test_auth_problem_pauses_instead_of_failing(client):
    p = client.post("/api/profiles", json={"name": "Outlook", "kind": "graph", "from_email": "x@y.test",
                                           "config": {"client_id": "00000000-0000-0000-0000-000000000000"}}).json()
    c = client.post("/api/campaigns", json={"name": "Needs login"}).json()
    client.put(f"/api/campaigns/{c['id']}", json={"profile_id": p["id"], "subject": "s", "body": "b"})
    client.post(f"/api/campaigns/{c['id']}/run", json={"mode": "draft"})
    c2 = wait_done(client, c["id"])
    assert c2["status"] == "paused" and c2["note"]
    assert c2["counts"].get("failed", 0) == 0 and c2["counts"]["pending"] >= 1


def test_csv_import_and_bad_file(client):
    csvb = "Name;E-mail\nZed;zed@example.com\n".encode()
    pv = client.post("/api/import/preview", files={"file": ("x.csv", csvb)}).json()
    assert pv["email_column"] == "E-mail" and pv["headers"] == ["Name", "E-mail"]
    assert client.post("/api/import/preview", files={"file": ("x.pdf", b"%PDF")}).status_code == 400
    assert client.post("/api/import/preview", files={"file": ("x.xlsx", b"garbage")}).status_code == 400


def test_password_never_returned(client):
    p = client.post("/api/profiles", json={"name": "S", "kind": "smtp", "from_email": "a@b.test",
                                           "config": {"host": "h", "password": "secret"}}).json()
    assert "password" not in p["config"] and p["has_password"]
    p2 = client.put(f"/api/profiles/{p['id']}", json={"name": "S2", "config": {"host": "h2", "password": ""}}).json()
    assert p2["has_password"]                                 # blank = keep
